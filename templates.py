# templates.py — OpenField Collect
# Template builder, questions, choices, imports
# SAFE for migrations, ordering, bulk import, and deletions

from __future__ import annotations

import re
from datetime import datetime
from typing import List, Dict, Tuple

from db import get_conn


# -------------------------------------------------
# Helpers
# -------------------------------------------------

def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _table_columns(table: str) -> List[str]:
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"PRAGMA table_info({table})")
        return [r["name"] for r in cur.fetchall()]


def _order_col_questions() -> str:
    cols = _table_columns("template_questions")
    if "display_order" in cols:
        return "display_order"
    if "order_no" in cols:
        return "order_no"
    return "display_order"


def _order_col_choices() -> str:
    cols = _table_columns("template_question_choices")
    if "display_order" in cols:
        return "display_order"
    if "order_no" in cols:
        return "order_no"
    return "display_order"


def _next_display_order(table: str, where_sql: str, params: tuple, order_col: str = "display_order") -> int:
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"SELECT COALESCE(MAX({order_col}), 0) + 1 AS n FROM {table} {where_sql}",
            params,
        )
        return int(cur.fetchone()["n"])


# -------------------------------------------------
# Templates
# -------------------------------------------------

def create_template(
    name: str,
    description: str = "",
    is_active: int = 1,
    require_enumerator_code: int = 0,
    enable_gps: int = 0,
    enable_coverage: int = 0,
    coverage_scheme_id: int | None = None,
    project_id: int | None = None,
    created_by: str | None = None,
    source: str = "manual",
    assignment_mode: str = "INHERIT",
    template_version: str = "v1",
    enable_consent: int = 0,
    enable_attestation: int = 0,
    is_sensitive: int = 0,
    restricted_exports: int = 0,
    redacted_fields: str | None = None,
) -> int:
    created_by = (created_by or "System").strip() or "System"
    cols = set(_table_columns("survey_templates"))
    now = _now()
    values_map = {
        "name": name.strip(),
        "description": description.strip(),
        "is_active": int(is_active),
        "created_at": now,
        "created_by": created_by,
        "updated_at": now,
        "source": (source or "manual").strip().lower(),
        "assignment_mode": (assignment_mode or "INHERIT").strip().upper(),
        "template_version": (template_version or "v1").strip(),
        "enable_consent": int(enable_consent),
        "enable_attestation": int(enable_attestation),
        "is_sensitive": int(is_sensitive),
        "restricted_exports": int(restricted_exports),
        "redacted_fields": (redacted_fields or "").strip() or None,
        "require_enumerator_code": int(require_enumerator_code),
        "enable_gps": int(enable_gps),
        "enable_coverage": int(enable_coverage),
        "coverage_scheme_id": coverage_scheme_id,
        "project_id": project_id,
    }
    insert_cols = [c for c in values_map.keys() if c in cols]
    if not insert_cols:
        raise RuntimeError("survey_templates schema is not initialized.")
    placeholders = ", ".join(["?"] * len(insert_cols))
    sql = f"INSERT INTO survey_templates ({', '.join(insert_cols)}) VALUES ({placeholders})"
    insert_values = tuple(values_map[c] for c in insert_cols)
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(sql, insert_values)
        conn.commit()
        return int(cur.lastrowid)


def list_templates(limit: int = 200, project_id: int | None = None) -> List[Tuple]:
    cols = set(_table_columns("survey_templates"))
    where_parts: List[str] = []
    params: List[int] = []
    if project_id is not None and "project_id" in cols:
        where_parts.append("project_id=?")
        params.append(int(project_id))
    if "deleted_at" in cols:
        where_parts.append("deleted_at IS NULL")
    where = ("WHERE " + " AND ".join(where_parts)) if where_parts else ""
    params.append(int(limit))
    created_by_sel = "created_by" if "created_by" in cols else "NULL AS created_by"
    updated_at_sel = "updated_at" if "updated_at" in cols else "NULL AS updated_at"
    source_sel = "source" if "source" in cols else "NULL AS source"
    assignment_mode_sel = "assignment_mode" if "assignment_mode" in cols else "'INHERIT' AS assignment_mode"
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT id, name, description, created_at, {created_by_sel}, {updated_at_sel}, {source_sel}, {assignment_mode_sel}
            FROM survey_templates
            {where}
            ORDER BY id DESC
            LIMIT ?
            """,
            tuple(params),
        )
        return cur.fetchall()


def get_template_config(template_id: int) -> Dict:
    cols = set(_table_columns("survey_templates"))
    where = "id=?"
    if "deleted_at" in cols:
        where += " AND deleted_at IS NULL"
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM survey_templates WHERE {where} LIMIT 1", (int(template_id),))
        r = cur.fetchone()
        return dict(r) if r else {}


def set_template_config(template_id: int, **kwargs) -> None:
    if not kwargs:
        return
    cols = _table_columns("survey_templates")
    fields = []
    values = []
    for k, v in kwargs.items():
        if k in cols:
            fields.append(f"{k}=?")
            values.append(v)
    if not fields:
        return
    if "updated_at" in cols:
        fields.append("updated_at=?")
        values.append(_now())
    where_sql = "id=?"
    if "deleted_at" in cols:
        where_sql += " AND deleted_at IS NULL"
    with get_conn() as conn:
        conn.execute(
            f"UPDATE survey_templates SET {', '.join(fields)} WHERE {where_sql}",
            (*values, int(template_id)),
        )
        conn.commit()


def soft_delete_template(template_id: int) -> None:
    cols = set(_table_columns("survey_templates"))
    with get_conn() as conn:
        if "deleted_at" in cols:
            if "updated_at" in cols:
                conn.execute(
                    "UPDATE survey_templates SET deleted_at=?, updated_at=? WHERE id=?",
                    (_now(), _now(), int(template_id)),
                )
            else:
                conn.execute(
                    "UPDATE survey_templates SET deleted_at=? WHERE id=?",
                    (_now(), int(template_id)),
                )
        else:
            conn.execute("DELETE FROM survey_templates WHERE id=?", (int(template_id),))
        conn.commit()


# -------------------------------------------------
# Questions
# -------------------------------------------------

def add_template_question(
    template_id: int,
    question_text: str,
    question_type: str = "TEXT",
    order_no: int | None = None,
    is_required: int = 0,
    help_text: str | None = None,
    validation_json: str | None = None,
) -> int:
    question_text = question_text.strip()
    if not question_text:
        raise ValueError("Question text is required")

    question_type = question_type.upper().strip()

    order_col = _order_col_questions()
    if order_no is None:
        order_no = _next_display_order(
            "template_questions",
            "WHERE template_id=?",
            (int(template_id),),
            order_col=order_col,
        )

    cols = _table_columns("template_questions")
    fields = ["template_id", "question_text", "question_type", order_col, "is_required", "created_at"]
    values = [
        int(template_id),
        question_text,
        question_type,
        int(order_no),
        int(is_required),
        _now(),
    ]
    if "help_text" in cols:
        fields.append("help_text")
        values.append((help_text or "").strip() or None)
    if "validation_json" in cols:
        fields.append("validation_json")
        values.append((validation_json or "").strip() or None)

    placeholders = ",".join(["?"] * len(fields))
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            INSERT INTO template_questions
              ({", ".join(fields)})
            VALUES ({placeholders})
            """,
            tuple(values),
        )
        conn.commit()
        return int(cur.lastrowid)


def get_template_questions(template_id: int) -> List[Tuple]:
    order_col = _order_col_questions()
    cols = _table_columns("template_questions")
    help_expr = "help_text" if "help_text" in cols else "NULL AS help_text"
    val_expr = "validation_json" if "validation_json" in cols else "NULL AS validation_json"
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT id, question_text, question_type, {order_col} AS display_order, is_required,
                   {help_expr}, {val_expr}
            FROM template_questions
            WHERE template_id=?
            ORDER BY display_order ASC, id ASC
            """,
            (int(template_id),),
        )
        return cur.fetchall()


# -------------------------------------------------
# Choices
# -------------------------------------------------

def add_choice(question_id: int, choice_text: str) -> int:
    choice_text = choice_text.strip()
    if not choice_text:
        raise ValueError("Choice text is required")

    order_col = _order_col_choices()
    order_no = _next_display_order(
        "template_question_choices",
        "WHERE template_question_id=?",
        (int(question_id),),
        order_col=order_col,
    )

    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            INSERT INTO template_question_choices
              (template_question_id, choice_text, {order_col}, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (
                int(question_id),
                choice_text,
                int(order_no),
                _now(),
            ),
        )
        conn.commit()
        return int(cur.lastrowid)


def list_choices(question_id: int) -> List[Tuple]:
    order_col = _order_col_choices()
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT id, template_question_id, choice_text, {order_col} AS display_order
            FROM template_question_choices
            WHERE template_question_id=?
            ORDER BY display_order ASC, id ASC
            """,
            (int(question_id),),
        )
        return cur.fetchall()


# -------------------------------------------------
# Import from TEXT / DOCX / PDF
# -------------------------------------------------
QUESTION_PREFIX_RE = re.compile(r"^\s*(?:Q?\d+[\)\.\:\-]|[A-Za-z][\)\.\:])\s+")
CHOICE_PREFIX_RE = re.compile(r"^\s*(?:[-\*]|\u2022|\d+[\)\.]|[A-Za-z][\)\.])\s+")
INLINE_SPLIT_RE = re.compile(r"\s*[,;/\|]\s*")
REQUIRED_MARK_RE = re.compile(r"(\*|\[required\]|\(required\))", re.IGNORECASE)


def _clean_leading_marker(text: str) -> str:
    s = (text or "").strip()
    s = re.sub(r"^\s*(?:Q?\d+[\)\.\:\-])\s+", "", s)
    s = re.sub(r"^\s*([A-Za-z][\)\.\:])\s+", "", s)
    s = re.sub(r"^\s*[-\*\u2022]\s+", "", s)
    return s.strip()


def _clean_choice(text: str) -> str:
    s = (text or "").strip()
    s = re.sub(r"^\s*(?:[-\*\u2022]|\d+[\)\.]|[A-Za-z][\)\.])\s+", "", s)
    return s.strip()


def _is_choice_line(text: str) -> bool:
    return bool(CHOICE_PREFIX_RE.match(text or ""))


def _is_question_line(text: str) -> bool:
    s = (text or "").strip()
    if not s:
        return False
    if s.endswith("?"):
        return True
    if ":" in s:
        return True
    if QUESTION_PREFIX_RE.match(s):
        return True
    if s.lower().startswith("q") and len(s) > 2 and s[1].isdigit():
        return True
    return False


def _strip_required_marker(text: str) -> Tuple[str, bool]:
    s = (text or "").strip()
    if not s:
        return s, False
    required = bool(REQUIRED_MARK_RE.search(s))
    s = REQUIRED_MARK_RE.sub("", s).strip()
    s = re.sub(r"\s+\*$", "", s).strip()
    return s, required


def _split_inline_choices(text: str) -> Tuple[str, List[str]]:
    s = (text or "").strip()
    if ":" not in s:
        return s, []
    left, right = s.split(":", 1)
    right = right.strip()
    if not right:
        return left.strip(), []

    # If right side looks like a list, use it as choices
    tokens = [t.strip() for t in INLINE_SPLIT_RE.split(right) if t.strip()]
    if len(tokens) >= 2:
        return left.strip(), tokens

    # If it looks like YES/NO but not split well
    if re.search(r"\byes\s*/\s*no\b", right, re.IGNORECASE):
        return left.strip(), ["Yes", "No"]

    # Otherwise treat as part of the question text
    if re.fullmatch(r"[_\-\.\s]+", right or ""):
        return left.strip(), []
    return f"{left.strip()} {right}".strip(), []


def _infer_question_type(question_text: str, choices: List[str]) -> str:
    q = (question_text or "").strip()
    lower = q.lower()

    multi_hint = any(
        k in lower
        for k in ["select all", "choose all", "all that apply", "multiple answers", "check all"]
    )

    if choices:
        norm = {c.strip().lower() for c in choices if c.strip()}
        if norm == {"yes", "no"}:
            return "YESNO"
        if multi_hint:
            return "MULTI_CHOICE"
        if len(choices) >= 7:
            return "DROPDOWN"
        return "SINGLE_CHOICE"

    if "yes/no" in lower or "yes or no" in lower:
        return "YESNO"
    if any(k in lower for k in ["email"]):
        return "EMAIL"
    if any(k in lower for k in ["phone", "mobile", "tel", "telephone"]):
        return "PHONE"
    if any(k in lower for k in ["date", "dob", "birth", "day of", "day"]):
        return "DATE"
    if any(k in lower for k in ["how many", "number of", "count", "quantity", "age", "minutes", "hours", "%", "rate"]):
        return "NUMBER"
    if any(k in lower for k in ["describe", "explain", "details", "why", "how", "comment", "observations", "notes"]):
        return "LONGTEXT"
    return "TEXT"


def parse_questions_from_text(raw_text: str, max_questions: int = 200) -> List[Dict]:
    """
    Parses raw text into question dicts:
      {question_text, question_type, choices, is_required}
    """
    text = raw_text or ""
    lines = [ln.rstrip() for ln in text.splitlines()]
    items: List[Dict] = []
    current: Dict | None = None

    def push_current() -> None:
        nonlocal current
        if not current:
            return
        q_text = (current.get("question_text") or "").strip()
        if not q_text:
            current = None
            return
        q_text, req_from_text = _strip_required_marker(q_text)
        required = 1 if current.get("is_required") or req_from_text else 0
        choices = [c for c in (current.get("choices") or []) if c]
        q_type = _infer_question_type(q_text, choices)
        items.append(
            {
                "question_text": q_text,
                "question_type": q_type,
                "choices": choices,
                "is_required": required,
            }
        )
        current = None

    for ln in lines:
        s = ln.strip()
        if not s:
            push_current()
            continue

        if current and _is_choice_line(s):
            choice = _clean_choice(s)
            if choice:
                current.setdefault("choices", []).append(choice)
            continue

        if _is_question_line(s) or current is None:
            if current:
                push_current()
            q_text, inline_choices = _split_inline_choices(s)
            q_text = _clean_leading_marker(q_text)
            current = {
                "question_text": q_text,
                "choices": inline_choices,
                "is_required": 0,
            }
            if REQUIRED_MARK_RE.search(s):
                current["is_required"] = 1
            continue

        if current:
            current["question_text"] = (current.get("question_text") or "").strip() + " " + s

    push_current()
    return items[:max_questions]


def import_questions_from_items(
    template_id: int,
    items: List[Dict],
    default_required: int = 0,
) -> Dict:
    added = 0
    errors: List[str] = []

    for item in items:
        try:
            q_text = (item.get("question_text") or "").strip()
            if not q_text:
                continue
            q_type = (item.get("question_type") or "TEXT").upper()
            choices = item.get("choices") or []
            is_required = 1 if int(default_required) == 1 else int(item.get("is_required") or 0)

            qid = add_template_question(
                template_id,
                q_text,
                question_type=q_type,
                is_required=is_required,
            )

            if q_type in ("SINGLE_CHOICE", "DROPDOWN", "MULTI_CHOICE") and choices:
                for c in choices:
                    ctext = (c or "").strip()
                    if ctext:
                        add_choice(qid, ctext)

            added += 1
        except Exception as e:
            errors.append(f"{item.get('question_text') or 'Question'}: {e}")

    return {"added": added, "errors": errors}


def import_questions_from_text(
    template_id: int,
    raw_text: str,
    default_required: int = 0,
) -> Dict:
    items = parse_questions_from_text(raw_text)
    return import_questions_from_items(template_id, items, default_required=default_required)


def _extract_text_from_docx(docx_path: str) -> str:
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx is not installed")
    doc = Document(docx_path)
    parts = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    return "\n".join(parts)


def _extract_text_from_pdf(pdf_path: str) -> str:
    try:
        import fitz  # PyMuPDF
    except Exception as e:
        raise RuntimeError("PyMuPDF is not installed") from e
    doc = fitz.open(pdf_path)
    parts = []
    for page in doc:
        t = (page.get_text("text") or "").strip()
        if t:
            parts.append(t)
    doc.close()
    return "\n".join(parts)


def extract_text_from_docx(docx_path: str) -> str:
    return _extract_text_from_docx(docx_path)


def extract_text_from_pdf(pdf_path: str) -> str:
    return _extract_text_from_pdf(pdf_path)


def preview_questions_from_docx(docx_path: str, max_questions: int = 200) -> List[Dict]:
    raw = _extract_text_from_docx(docx_path)
    return parse_questions_from_text(raw, max_questions=max_questions)


def preview_questions_from_pdf(pdf_path: str, max_questions: int = 200) -> List[Dict]:
    raw = _extract_text_from_pdf(pdf_path)
    return parse_questions_from_text(raw, max_questions=max_questions)


def import_questions_from_docx(
    template_id: int,
    docx_path: str,
    default_required: int = 0,
) -> Dict:
    raw = _extract_text_from_docx(docx_path)
    items = parse_questions_from_text(raw)
    return import_questions_from_items(template_id, items, default_required=default_required)


def import_questions_from_pdf(
    template_id: int,
    pdf_path: str,
    default_required: int = 0,
) -> Dict:
    raw = _extract_text_from_pdf(pdf_path)
    items = parse_questions_from_text(raw)
    return import_questions_from_items(template_id, items, default_required=default_required)
